"""
Data Converter: Excel → JSON
Mengkonversi Laporan Keuangan Karim Gadget dari Excel ke JSON
Versi 2.0 — Include Data Penjualan
"""
import openpyxl
import json
import uuid
import os
from datetime import datetime

def convert_excel_to_json():
    wb = openpyxl.load_workbook('Laporan Keuangan Karim Gadget ID.xlsx', data_only=True)
    
    # Mapping sheet names to month/year
    sheet_month_map = {
        'Ag-Sep': {'start_month': 8, 'start_year': 2025},
        'Okt': {'start_month': 10, 'start_year': 2025},
        'Nov': {'start_month': 11, 'start_year': 2025},
        'Des': {'start_month': 12, 'start_year': 2025},
        'Jan': {'start_month': 1, 'start_year': 2026},
        'Feb': {'start_month': 2, 'start_year': 2026},
        'Mar': {'start_month': 3, 'start_year': 2026},
        'Apr': {'start_month': 4, 'start_year': 2026},
        'Mei': {'start_month': 5, 'start_year': 2026},
        'Juni': {'start_month': 6, 'start_year': 2026},
        'Juli': {'start_month': 7, 'start_year': 2026},
        'Agu': {'start_month': 8, 'start_year': 2026},
    }
    
    all_transactions = []
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        meta = sheet_month_map.get(sheet_name, {})
        last_date = None
        
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
            # Skip empty rows
            if all(v is None for v in row[:7]):
                continue
            
            tanggal_raw = row[0]
            deskripsi = row[1]
            jumlah = row[2]
            kategori = row[3]
            uang_masuk = row[4]
            uang_keluar = row[5]
            saldo = row[6]
            
            # Skip if no meaningful data
            if deskripsi is None and uang_masuk is None and uang_keluar is None:
                continue
            
            # Handle date
            if tanggal_raw is not None:
                if isinstance(tanggal_raw, datetime):
                    last_date = tanggal_raw.strftime('%Y-%m-%d')
                elif isinstance(tanggal_raw, str):
                    import re
                    m = re.match(r'^(\d{1,2})[/\-](\d{1,2})[/\-]?(\d{4})$', tanggal_raw)
                    if m:
                        d, m_, y = m.groups()
                        last_date = f"{y}-{int(m_):02d}-{int(d):02d}"
                    else:
                        last_date = tanggal_raw
            
            # Use last known date for rows without date (merged cells)
            tanggal = last_date
            
            # Normalize jumlah
            if jumlah == '-' or jumlah is None:
                jumlah_val = 0
            else:
                try:
                    jumlah_val = int(jumlah)
                except (ValueError, TypeError):
                    jumlah_val = 0
            
            # Normalize kategori
            if kategori:
                kategori = kategori.strip()
            else:
                kategori = 'Lainnya'
            
            # Normalize money values
            uang_masuk_val = float(uang_masuk) if uang_masuk else 0
            uang_keluar_val = float(uang_keluar) if uang_keluar else 0
            saldo_val = float(saldo) if saldo else 0
            
            transaction = {
                'id': str(uuid.uuid4())[:8],
                'tanggal': tanggal,
                'deskripsi': str(deskripsi) if deskripsi else '',
                'jumlah': jumlah_val,
                'kategori': kategori,
                'uangMasuk': uang_masuk_val,
                'uangKeluar': uang_keluar_val,
                'saldo': saldo_val,
                'bulan': sheet_name
            }
            
            all_transactions.append(transaction)
    
    # Sort by date
    all_transactions.sort(key=lambda x: x['tanggal'] or '0000-00-00')
    
    # ============ Baca Data Penjualan ============
    sales_data = []
    try:
        wb_sales = openpyxl.load_workbook('Data penjualan karim Gadget.xlsx', data_only=True)
        if 'Real Time' in wb_sales.sheetnames:
            ws_sales = wb_sales['Real Time']
        else:
            ws_sales = wb_sales.active
        
        print(f"\n[*] Memproses Data Penjualan...")
        
        for row_idx, row in enumerate(ws_sales.iter_rows(min_row=2, values_only=True), 2):
            # Skip empty rows
            if all(v is None for v in row[:8]):
                continue
            
            # Kolom: Nota, Tanggal Masuk, Tanggal Keluar, Tipe/Model, Harga Beli, Harga Jual, Profit, Keterangan
            nota = row[0]
            tanggal_masuk_raw = row[1]
            tanggal_keluar_raw = row[2]
            tipe_model = row[3]
            harga_beli = row[4]
            harga_jual = row[5]
            profit = row[6]
            keterangan = row[7] if len(row) > 7 else None
            
            # Skip jika tidak ada data penting
            if tipe_model is None and harga_jual is None:
                continue
            
            # Handle tanggal masuk
            tanggal_masuk = None
            if tanggal_masuk_raw is not None:
                if isinstance(tanggal_masuk_raw, datetime):
                    tanggal_masuk = tanggal_masuk_raw.strftime('%Y-%m-%d')
                elif isinstance(tanggal_masuk_raw, str):
                    tanggal_masuk = tanggal_masuk_raw
                elif isinstance(tanggal_masuk_raw, (int, float)):
                    # Excel date serial
                    try:
                        from datetime import date
                        epoch = datetime(1899, 12, 30)
                        tanggal_masuk = (epoch + __import__('datetime').timedelta(days=int(tanggal_masuk_raw))).strftime('%Y-%m-%d')
                    except:
                        tanggal_masuk = str(tanggal_masuk_raw)
            
            # Handle tanggal keluar
            tanggal_keluar = None
            if tanggal_keluar_raw is not None:
                if isinstance(tanggal_keluar_raw, datetime):
                    tanggal_keluar = tanggal_keluar_raw.strftime('%Y-%m-%d')
                elif isinstance(tanggal_keluar_raw, str):
                    tanggal_keluar = tanggal_keluar_raw
                elif isinstance(tanggal_keluar_raw, (int, float)):
                    try:
                        epoch = datetime(1899, 12, 30)
                        tanggal_keluar = (epoch + __import__('datetime').timedelta(days=int(tanggal_keluar_raw))).strftime('%Y-%m-%d')
                    except:
                        tanggal_keluar = str(tanggal_keluar_raw)
            
            # Normalisasi tahun salah (2020, 2024 -> 2026)
            if tanggal_keluar and tanggal_keluar.startswith('2020-'):
                tanggal_keluar = '2026-' + tanggal_keluar[5:]
            if tanggal_keluar and tanggal_keluar.startswith('2024-'):
                tanggal_keluar = '2026-' + tanggal_keluar[5:]
            if tanggal_masuk and tanggal_masuk.startswith('2024-'):
                tanggal_masuk = '2026-' + tanggal_masuk[5:]

            # Hitung turnover (waktu putar dalam hari)
            turnover_days = None
            if tanggal_masuk and tanggal_keluar:
                try:
                    dt_masuk = datetime.strptime(tanggal_masuk, '%Y-%m-%d')
                    dt_keluar = datetime.strptime(tanggal_keluar, '%Y-%m-%d')
                    turnover_days = (dt_keluar - dt_masuk).days
                except:
                    turnover_days = None
            
            # Normalize values
            harga_beli_val = float(harga_beli) if harga_beli is not None else 0
            harga_jual_val = float(harga_jual) if harga_jual is not None else 0
            profit_val = float(profit) if profit is not None else (harga_jual_val - harga_beli_val)
            
            sale = {
                'id': str(uuid.uuid4())[:8],
                'nota': str(nota) if nota else f'NOTA-{row_idx}',
                'tanggalMasuk': tanggal_masuk,
                'tanggalKeluar': tanggal_keluar,
                'tipeModel': str(tipe_model) if tipe_model else '',
                'hargaBeli': harga_beli_val,
                'hargaJual': harga_jual_val,
                'profit': profit_val,
                'keterangan': str(keterangan) if keterangan else '',
                'turnoverDays': turnover_days,
            }
            
            sales_data.append(sale)
        
        # Sort by tanggal keluar
        sales_data.sort(key=lambda x: x['tanggalKeluar'] or '0000-00-00')
        
        print(f"[OK] Berhasil membaca {len(sales_data)} data penjualan")
        
    except FileNotFoundError:
        print("[!] File 'Data penjualan karim Gadget.xlsx' tidak ditemukan, skip sales data")
    except Exception as e:
        print(f"[!] Error membaca data penjualan: {e}")
    
    # ============ Output ============
    # Ensure output directory exists
    os.makedirs('data', exist_ok=True)
    
    # Write transactions JSON
    with open('data/transactions.json', 'w', encoding='utf-8') as f:
        json.dump(all_transactions, f, ensure_ascii=False, indent=2)
    
    # Write sales JSON
    with open('data/sales.json', 'w', encoding='utf-8') as f:
        json.dump(sales_data, f, ensure_ascii=False, indent=2)
    
    # Write combined initial-data.js (embedded for file:// protocol)
    with open('data/initial-data.js', 'w', encoding='utf-8') as f:
        f.write('// Auto-generated by data-converter.py — DO NOT EDIT MANUALLY\n')
        f.write(f'// Generated: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}\n\n')
        f.write('const INITIAL_DATA = ')
        json.dump({'transactions': all_transactions, 'sales': sales_data}, f, ensure_ascii=False, indent=2)
        f.write(';\n')
    
    print(f"\n[OK] Berhasil mengkonversi {len(all_transactions)} transaksi")
    print(f"[OK] Berhasil memproses {len(sales_data)} data penjualan")
    print(f"[>>] Output: data/transactions.json, data/sales.json, data/initial-data.js")
    
    # Print summary per month
    from collections import Counter
    month_count = Counter(t['bulan'] for t in all_transactions)
    print("\n[*] Ringkasan per bulan:")
    for sheet_name in wb.sheetnames:
        count = month_count.get(sheet_name, 0)
        print(f"   {sheet_name}: {count} transaksi")
    
    # Print category summary
    cat_count = Counter(t['kategori'] for t in all_transactions)
    print("\n[*] Ringkasan per kategori:")
    for cat, count in cat_count.most_common():
        print(f"   {cat}: {count} transaksi")

if __name__ == '__main__':
    convert_excel_to_json()
